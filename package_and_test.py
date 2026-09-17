"""Create the shareable ZIP and test the executable from a fresh extraction."""
import hashlib
import json
import os
from pathlib import Path
import subprocess
import tempfile
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
source = ROOT/'dist'/'格识'
release = ROOT/'release'; release.mkdir(exist_ok=True)
archive = release/'格识-1.0.0-Windows-x64-便携版.zip'
with ZipFile(archive, 'w', ZIP_DEFLATED, compresslevel=6) as output:
    for file in sorted(source.rglob('*')):
        if file.is_file():
            output.write(file, Path('格识')/file.relative_to(source))
with ZipFile(archive) as compressed:
    assert compressed.testzip() is None
    assert any(name.endswith('.onnx') for name in compressed.namelist())
    assert not any('/.venv/' in name or '/test-output/' in name or name.endswith('.log') for name in compressed.namelist())
    destination = Path(tempfile.mkdtemp(prefix='gezhi-portable-'))/'新目录 中文 space'
    compressed.extractall(destination)
app = destination/'格识'/'格识.exe'
output = ROOT/'test-output'/'portable-exe'
env = {key:value for key,value in os.environ.items()
       if key.upper() not in ['PYTHONHOME','PYTHONPATH','VIRTUAL_ENV','TCL_LIBRARY','TK_LIBRARY','CONDA_PREFIX']}
env['PATH'] = str(Path(os.environ['WINDIR'])/'System32')
env['PYTHONNOUSERSITE'] = '1'
result = subprocess.run([str(app), '--self-test', str(destination/'格识'/'示例图片'/'中文表格示例.png'), str(output)],
                        cwd=destination, env=env, timeout=120, creationflags=subprocess.CREATE_NO_WINDOW)
assert result.returncode == 0, f'Frozen app failed with exit {result.returncode}; inspect %LOCALAPPDATA%/GezhiOCR/desktop.log'
report = json.loads((output/'result.json').read_text(encoding='utf-8'))
assert report['frozen'] and report['status']=='ok' and report['network']=='blocked'
book = load_workbook(output/'独立软件测试.xlsx')
assert book.active['B2'].value == '独立软件中校对'
assert book.active['C2'].value == '00123' and book.active['C2'].data_type == 's'
assert book.active.max_row == 4 and book.active.max_column == 3
checksum = hashlib.file_digest(archive.open('rb'), 'sha256').hexdigest()
(release/(archive.name+'.sha256')).write_text(f'{checksum}  {archive.name}\n',encoding='utf-8')
summary = {'archive':str(archive),'size_bytes':archive.stat().st_size,'sha256':checksum,
           'extracted_exe':str(app),'test':'passed','offline_ocr':'passed',
           'gui_edit':'passed','xlsx_readback':'passed','no_python_on_path':True,
           'process_exited':True}
(ROOT/'test-output'/'portable-validation.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False,indent=2))
